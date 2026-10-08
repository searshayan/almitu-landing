#!/usr/bin/env python3
"""Convert the approved Speaking Curriculum workbook into app/curriculum/speaking-<tier>.json.

Usage: python3 tools/build_speaking_curriculum.py "/path/to/Speaking Curriculum ... (Completed Draft).xlsx"

The workbook is the source of truth. The hidden "Difficulty Step" column is carried
through as `step` (generator only; never shown to tutors or learners).
"""
import json, re, sys, os
import openpyxl

SHEETS = {
    'Speaking - Foundation Tier': ('foundation', 'Speaking Foundation'),
    'Speaking - Development Tier': ('development', 'Speaking Development'),
    'Speaking - Proficiency Tier': ('proficiency', 'Speaking Proficiency'),
}
# Transactional sessions: Listening is a short setup message (voicemail, announcement)
# instead of a model talk, because a monologue about "what I said at the bank" would not
# model the learner's own lines.
SETUP_TITLE = re.compile(r"^(at the |taking a taxi|renting a car|buying train tickets|returning an item|"
                         r"shopping for clothes|asking for directions|online shopping|job ads)", re.I)

def lines(v):
    return [s.strip() for s in str(v or '').split('\n') if s.strip()]

def build(xlsx, outdir):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    for sheet, (tier, level) in SHEETS.items():
        ws = wb[sheet]
        head = {str(c.value).strip(): i for i, c in enumerate(ws[2]) if c.value}
        col = lambda n: next(head[h] for h in head if h.startswith(n))
        idx = {
            'id': head['Session ID'], 'title': head['Session Title'], 'obj': head['Learning Objective'],
            'unit': col('Unit'), 'outcome': col('Unit Outcome'), 'type': col('Session Type'),
            'cando': col('Can-Do'), 'expr': col('Target Expressions'), 'lang': col('Language Focus'),
            'pron': col('Pronunciation'), 'act': col('Speaking Activity'), 'task': col('Speaking Task'),
            'chal': col('Between-Session'), 'exit': col('Exit Task'), 'notes': col('Tutor Notes'),
            'repair': col('Repair Phrases'), 'step': col('Difficulty Step'),
        }
        recs = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            sid = row[idx['id']]
            if not sid: continue
            n = int(str(sid).split('-')[-1])
            um = re.match(r'Unit (\d+): (.*)', str(row[idx['unit']]))
            kind = 'review' if str(row[idx['type']]).startswith('Review') else 'core'
            title = str(row[idx['title']]).strip()
            rec = {
                'curriculum_id': sid, 'skill': 'speaking', 'level': level, 'tier': tier,
                'step': row[idx['step']], 'unit': int(um.group(1)), 'unit_title': um.group(2),
                'unit_outcome': row[idx['outcome']], 'kind': kind, 'title': title,
                'objective': row[idx['obj']], 'can_do': row[idx['cando']],
                'expressions': lines(row[idx['expr']]), 'language_focus': row[idx['lang']],
                'pronunciation': row[idx['pron']], 'activity': row[idx['act']], 'task': row[idx['task']],
                'exit_task': row[idx['exit']], 'tutor_notes': row[idx['notes']],
                'challenge': [re.sub(r'^\d+\.\s*', '', s) for s in lines(row[idx['chal']])],
                'repair_phrases': lines(row[idx['repair']]) if kind == 'core' else [],
                'listening_mode': 'setup_message' if (kind == 'core' and SETUP_TITLE.match(title)) else 'model_talk',
            }
            if kind == 'review':
                m = re.search(r'Lessons (\d+)\D+(\d+)', str(row[idx['cando']]) + ' ' + str(row[idx['pron']]))
                lo, hi = int(m.group(1)), int(m.group(2))
                rec['covers'] = ['S-%s-%03d' % (tier.capitalize(), k) for k in range(lo, hi + 1)]
            recs.append(rec)
        # validation
        ids = [r['curriculum_id'] for r in recs]
        assert len(ids) == len(set(ids)), 'duplicate ids in ' + sheet
        for r in recs:
            want = 6 if r['step'] in ('F1', 'F2', 'F3', 'F4', 'D1', 'D2', 'D3') else 8
            assert len(r['expressions']) == want, (r['curriculum_id'], len(r['expressions']), want)
            assert len(r['challenge']) == 3, r['curriculum_id']
            blob = json.dumps(r)
            assert not re.search(r'\b(CEFR|Pre-A1|A1|A2|B1|B2|C1|C2)\b', blob), ('level label in ' + r['curriculum_id'])
        path = os.path.join(outdir, 'speaking-%s.json' % tier)
        with open(path, 'w', encoding='utf-8') as f:
            json.dump({'level': level, 'tier': 'speaking', 'sessions': recs}, f, ensure_ascii=False, indent=1)
        print(path, len(recs), 'sessions')

if __name__ == '__main__':
    root = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'app', 'curriculum')
    build(sys.argv[1], root)
